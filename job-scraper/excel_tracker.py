"""Sync saved jobs from SQLite to a human-maintained Excel tracker."""

import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


DATABASE_PATH = Path(__file__).resolve().with_name("jobs.db")
EXCEL_PATH = Path(__file__).resolve().with_name("job_tracker.xlsx")
STATUS_OPTIONS = (
    "Not Applied",
    "Applied",
    "OA",
    "Interview",
    "Rejected",
    "Offer",
    "Withdrawn",
)
HEADERS = (
    "Company",
    "Job Title",
    "Location",
    "Posted Date",
    "AI Score",
    "AI Match",
    "URL",
    "Status",
    "Applied Date",
    "Interview",
    "Notes",
    "AI Reason",
    "Matched Requirements",
    "Partial Matches",
    "Missing Requirements",
    "Local AI Score",
    "Local AI Match",
    "Local AI Reason",
)
MANUAL_HEADERS = ("Status", "Applied Date", "Interview", "Notes")
DATABASE_COLUMNS = (
    "company",
    "title",
    "location",
    "posted_at",
    "ai_score",
    "ai_match",
    "url",
    "ai_reason",
    "matched_requirements",
    "partial_matches",
    "missing_requirements",
    "local_ai_score",
    "local_ai_relevant",
    "local_ai_reason",
)
HEADER_TO_DATABASE = {
    "Company": "company",
    "Job Title": "title",
    "Location": "location",
    "Posted Date": "posted_at",
    "AI Score": "ai_score",
    "AI Match": "ai_match",
    "URL": "url",
    "AI Reason": "ai_reason",
    "Matched Requirements": "matched_requirements",
    "Partial Matches": "partial_matches",
    "Missing Requirements": "missing_requirements",
    "Local AI Score": "local_ai_score",
    "Local AI Match": "local_ai_relevant",
    "Local AI Reason": "local_ai_reason",
}
DEFAULT_WIDTHS = {
    "Company": 20,
    "Job Title": 36,
    "Location": 24,
    "Posted Date": 18,
    "AI Score": 12,
    "AI Match": 12,
    "URL": 48,
    "Status": 18,
    "Applied Date": 18,
    "Interview": 24,
    "Notes": 40,
    "AI Reason": 42,
    "Matched Requirements": 42,
    "Partial Matches": 42,
    "Missing Requirements": 42,
    "Local AI Score": 16,
    "Local AI Match": 16,
    "Local AI Reason": 42,
}


def _read_jobs(db_path):
    database_uri = f"{Path(db_path).resolve().as_uri()}?mode=ro"
    try:
        connection = sqlite3.connect(database_uri, uri=True)
        connection.row_factory = sqlite3.Row
        try:
            return [dict(row) for row in connection.execute(
                f"SELECT {', '.join(DATABASE_COLUMNS)} FROM jobs ORDER BY id"
            )]
        finally:
            connection.close()
    except sqlite3.Error as error:
        raise RuntimeError(f"Could not read jobs from {db_path}: {error}") from error


def _display_value(column, value):
    if value is None:
        return None
    if column in ("ai_match", "local_ai_relevant"):
        return bool(value)
    if column in ("matched_requirements", "partial_matches", "missing_requirements"):
        try:
            decoded = json.loads(value)
        except (TypeError, json.JSONDecodeError):
            return value
        if isinstance(decoded, list):
            return "\n".join(decoded)
    return value


def _column_indexes(worksheet):
    indexes = {
        cell.value: cell.column
        for cell in worksheet[1]
        if isinstance(cell.value, str) and cell.value
    }
    next_column = worksheet.max_column + 1
    for header in HEADERS:
        if header not in indexes:
            worksheet.cell(row=1, column=next_column, value=header)
            indexes[header] = next_column
            next_column += 1
    return indexes


def _merge_duplicate_rows(worksheet, indexes):
    url_column = indexes["URL"]
    manual_columns = [indexes[header] for header in MANUAL_HEADERS]
    first_row_by_url = {}
    rows_to_delete = []

    for row_number in range(2, worksheet.max_row + 1):
        url = worksheet.cell(row=row_number, column=url_column).value
        if not isinstance(url, str) or not url.strip():
            continue
        url = url.strip()
        if url not in first_row_by_url:
            first_row_by_url[url] = row_number
            continue

        first_row = first_row_by_url[url]
        for column in manual_columns:
            kept = worksheet.cell(row=first_row, column=column)
            duplicate = worksheet.cell(row=row_number, column=column)
            if kept.value in (None, "") and duplicate.value not in (None, ""):
                kept.value = duplicate.value
        rows_to_delete.append(row_number)

    for row_number in reversed(rows_to_delete):
        worksheet.delete_rows(row_number)

    first_row_by_url = {}
    for row_number in range(2, worksheet.max_row + 1):
        url = worksheet.cell(row=row_number, column=url_column).value
        if isinstance(url, str) and url.strip():
            first_row_by_url.setdefault(url.strip(), row_number)
    return first_row_by_url


def _format_worksheet(worksheet, indexes):
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(max(indexes.values()))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(worksheet.max_row, 1)}"
    worksheet.row_dimensions[1].height = 30

    for header, column in indexes.items():
        worksheet.column_dimensions[get_column_letter(column)].width = DEFAULT_WIDTHS.get(
            header, 20
        )

    status_column = get_column_letter(indexes["Status"])
    status_range = f"{status_column}2:{status_column}1048576"
    status_formula = '"' + ",".join(STATUS_OPTIONS) + '"'
    if not any(
        validation.formula1 == status_formula and status_range in str(validation.sqref)
        for validation in worksheet.data_validations.dataValidation
    ):
        validation = DataValidation(type="list", formula1=status_formula, allow_blank=True)
        validation.error = "Choose a status from the dropdown list."
        validation.errorTitle = "Invalid status"
        validation.prompt = "Select the current application status."
        worksheet.add_data_validation(validation)
        validation.add(status_range)


def sync_excel(db_path=DATABASE_PATH, excel_path=EXCEL_PATH):
    """Update database-owned spreadsheet fields while retaining manual tracking data."""
    jobs = _read_jobs(db_path)
    excel_path = Path(excel_path)
    if excel_path.exists():
        workbook = load_workbook(excel_path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Jobs"
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=1, column=column, value=header)

    indexes = _column_indexes(worksheet)
    existing_rows = _merge_duplicate_rows(worksheet, indexes)
    added = 0
    updated = 0

    for job in jobs:
        url = job.get("url")
        if not isinstance(url, str) or not url.strip():
            continue
        url = url.strip()
        row_number = existing_rows.get(url)
        if row_number is None:
            row_number = worksheet.max_row + 1
            existing_rows[url] = row_number
            worksheet.cell(row=row_number, column=indexes["Status"], value="Not Applied")
            added += 1
        else:
            updated += 1

        for header, database_column in HEADER_TO_DATABASE.items():
            if header in MANUAL_HEADERS:
                continue
            value = _display_value(database_column, job.get(database_column))
            cell = worksheet.cell(row=row_number, column=indexes[header])
            cell.value = value
            if header == "URL" and value:
                cell.hyperlink = value
                cell.style = "Hyperlink"

    _format_worksheet(worksheet, indexes)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(excel_path)
    return {"added": added, "updated": updated, "total": len(jobs), "path": excel_path}


def main():
    try:
        result = sync_excel()
    except (OSError, RuntimeError, sqlite3.Error) as error:
        print(f"Excel sync failed: {error}")
        return 1
    print(
        f"Excel sync complete: {result['added']} added, {result['updated']} updated; "
        f"{result['total']} database jobs. File: {result['path']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())